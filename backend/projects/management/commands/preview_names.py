import re

from openpyxl import load_workbook
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand


# ============================================================
# CONFIGURATION
# ============================================================

EXCEL_PATH = r"D:\shrunked.xlsx"

# None = use the active worksheet
SHEET_NAME = None

# Maximum number of Excel data rows to process
LIMIT = 50

# Names containing any of these values will be ignored
IGNORE_NAMES = [
    "kifuko-koech",
]

EMAIL_DOMAIN = "kalro.org"

PASSWORD = "P@$$w0rd254"

# Excel columns containing names
NAME_COLUMNS = [
    "PI",
    "Co PI 1",
    "COPI 2",
]

# IMPORTANT:
# True  = preview only, no users created
# False = actually create users
DRY_RUN = False


# Values that should not be treated as names
INVALID_VALUES = {
    "",
    "-",
    "--",
    "n/a",
    "na",
    "none",
    "null",
    "nan",
    "nat",
    "nil",
}


# ============================================================
# MANAGEMENT COMMAND
# ============================================================

class Command(BaseCommand):

    help = "Preview or import PI and Co-PI users from an Excel file"

    def handle(self, *args, **options):

        User = get_user_model()

        # ====================================================
        # LOAD EXCEL
        # ====================================================

        self.stdout.write(
            self.style.NOTICE(
                f"Loading Excel file: {EXCEL_PATH}"
            )
        )

        try:
            wb = load_workbook(
                EXCEL_PATH,
                data_only=True,
            )

        except FileNotFoundError:
            self.stdout.write(
                self.style.ERROR(
                    f"Excel file not found: {EXCEL_PATH}"
                )
            )
            return

        except Exception as e:
            self.stdout.write(
                self.style.ERROR(
                    f"Failed to load Excel file: {e}"
                )
            )
            return

        # ====================================================
        # SELECT SHEET
        # ====================================================

        try:
            ws = (
                wb[SHEET_NAME]
                if SHEET_NAME
                else wb.active
            )

        except KeyError:
            self.stdout.write(
                self.style.ERROR(
                    f"Sheet not found: {SHEET_NAME}"
                )
            )
            return

        self.stdout.write(
            self.style.NOTICE(
                f"Using sheet: {ws.title}"
            )
        )

        # ====================================================
        # READ ROWS
        # ====================================================

        rows = list(
            ws.iter_rows(
                values_only=True
            )
        )

        if not rows:
            self.stdout.write(
                self.style.ERROR(
                    "The Excel sheet is empty."
                )
            )
            return

        # First row = headers
        header = [
            str(h).strip()
            if h is not None
            else ""
            for h in rows[0]
        ]

        # Apply limit
        data_rows = rows[
            1:1 + LIMIT
        ]

        # ====================================================
        # FIND REQUIRED COLUMNS
        # ====================================================

        col_idx = {}

        for name in NAME_COLUMNS:

            col_idx[name] = None

            for i, h in enumerate(header):

                if (
                    h.strip().upper()
                    == name.strip().upper()
                ):
                    col_idx[name] = i
                    break

        self.stdout.write("\nColumn indexes:")

        for key, value in col_idx.items():

            self.stdout.write(
                f"  {key!r}: {value}"
            )

        # ====================================================
        # CHECK MISSING COLUMNS
        # ====================================================

        missing = [
            key
            for key, value in col_idx.items()
            if value is None
        ]

        if missing:

            self.stdout.write(
                self.style.ERROR(
                    f"Missing columns: {missing}"
                )
            )

            self.stdout.write(
                f"Available columns: {header}"
            )

            return

        # ====================================================
        # HELPER FUNCTIONS
        # ====================================================

        def clean_value(value):

            if value is None:
                return ""

            text = str(value).strip()

            if text.lower() in INVALID_VALUES:
                return ""

            return text

        def strip_title(name):

            return re.sub(
                r"^(dr|prof|mr|mrs|ms|miss|eng|sir)\.?\s*",
                "",
                name,
                flags=re.IGNORECASE,
            )

        def split_name(full_name):

            name = clean_value(full_name)

            if not name:
                return "", ""

            # Remove text inside parentheses
            name = re.sub(
                r"\(.*?\)",
                " ",
                name,
            )

            # Remove titles
            name = strip_title(name)

            # Convert comma to space
            name = name.replace(
                ",",
                " ",
            )

            parts = [
                part
                for part in name.split()
                if part
            ]

            if not parts:
                return "", ""

            if len(parts) == 1:
                return parts[0], ""

            return (
                parts[0],
                parts[-1],
            )

        def is_ignored(full_name):

            low = full_name.lower()

            return any(
                bad.lower() in low
                for bad in IGNORE_NAMES
            )

        def clean_part(part):

            part = part.lower()

            # Remove apostrophes and other
            # characters unsuitable for usernames
            part = re.sub(
                r"[^a-z0-9-]",
                "",
                part,
            )

            return part

        def make_username(first, last):

            first_clean = clean_part(first)
            last_clean = clean_part(last)

            if not first_clean and not last_clean:
                return ""

            if not last_clean:
                return first_clean

            return (
                f"{first_clean}.{last_clean}"
            )

        # ====================================================
        # COUNTERS
        # ====================================================

        created = 0
        existing = 0
        skipped = 0
        failed = 0

        blank_count = 0
        duplicate_count = 0
        ignored_count = 0
        invalid_count = 0

        # Track names encountered in Excel
        seen_names = set()

        # ====================================================
        # TABLE HEADER
        # ====================================================

        self.stdout.write("")

        self.stdout.write(
            f"{'ROW':<5} "
            f"{'SOURCE COL':<12} "
            f"{'FULL NAME':<45} "
            f"{'FIRST':<15} "
            f"{'LAST':<15} "
            f"{'USERNAME':<30} "
            f"{'EMAIL':<35}"
        )

        self.stdout.write(
            "-" * 165
        )

        # ====================================================
        # PROCESS ROWS
        # ====================================================

        for row_number, row in enumerate(
            data_rows,
            start=2,
        ):

            for column_name, column_index in col_idx.items():

                try:

                    # ----------------------------------------
                    # Get name from Excel
                    # ----------------------------------------

                    if column_index >= len(row):
                        full_name = ""

                    else:
                        full_name = clean_value(
                            row[column_index]
                        )

                    # ----------------------------------------
                    # Blank value
                    # ----------------------------------------

                    if not full_name:

                        blank_count += 1
                        skipped += 1

                        continue

                    # ----------------------------------------
                    # Invalid value
                    # ----------------------------------------

                    if full_name.lower() in INVALID_VALUES:

                        invalid_count += 1
                        skipped += 1

                        continue

                    # ----------------------------------------
                    # Ignored name
                    # ----------------------------------------

                    if is_ignored(full_name):

                        ignored_count += 1
                        skipped += 1

                        continue

                    # ----------------------------------------
                    # Duplicate name
                    # ----------------------------------------

                    name_key = full_name.lower()

                    if name_key in seen_names:

                        duplicate_count += 1
                        skipped += 1

                        continue

                    seen_names.add(name_key)

                    # ----------------------------------------
                    # Split name
                    # ----------------------------------------

                    first, last = split_name(
                        full_name
                    )

                    # Make sure we actually have
                    # a usable name
                    if not first:

                        invalid_count += 1
                        skipped += 1

                        continue

                    # ----------------------------------------
                    # Generate username
                    # ----------------------------------------

                    username = make_username(
                        first,
                        last,
                    )

                    if not username:

                        invalid_count += 1
                        skipped += 1

                        continue

                    # ----------------------------------------
                    # Generate email
                    # ----------------------------------------

                    email = (
                        f"{username}@{EMAIL_DOMAIN}"
                    )

                    # ----------------------------------------
                    # Check whether user already exists
                    # ----------------------------------------

                    user_exists = User.objects.filter(
                        email__iexact=email
                    ).exists()

                    if user_exists:

                        existing += 1

                        self.stdout.write(
                            f"{row_number:<5} "
                            f"{column_name:<12} "
                            f"{full_name!r:<45} "
                            f"{first!r:<15} "
                            f"{last!r:<15} "
                            f"{username:<30} "
                            f"{email:<35} "
                            "[EXISTS]"
                        )

                        continue

                    # ----------------------------------------
                    # Make username unique
                    # ----------------------------------------

                    original_username = username

                    counter = 1

                    while User.objects.filter(
                        username__iexact=username
                    ).exists():

                        counter += 1

                        username = (
                            f"{original_username}{counter}"
                        )

                    # Email follows final username
                    email = (
                        f"{username}@{EMAIL_DOMAIN}"
                    )

                    # ----------------------------------------
                    # Preview
                    # ----------------------------------------

                    self.stdout.write(
                        f"{row_number:<5} "
                        f"{column_name:<12} "
                        f"{full_name!r:<45} "
                        f"{first!r:<15} "
                        f"{last!r:<15} "
                        f"{username:<30} "
                        f"{email:<35}"
                    )

                    # ----------------------------------------
                    # DRY RUN
                    # ----------------------------------------

                    if DRY_RUN:
                        continue

                    # ----------------------------------------
                    # CREATE USER
                    # ----------------------------------------

                    User.objects.create_user(
                        username=username,
                        email=email,
                        password=PASSWORD,
                        first_name=first,
                        last_name=last,
                        is_staff=True,
                    )

                    created += 1

                except Exception as e:

                    failed += 1

                    self.stdout.write(
                        self.style.ERROR(
                            f"Row {row_number} "
                            f"[{column_name}]: "
                            f"FAILED -> {e}"
                        )
                    )

        # ====================================================
        # SUMMARY
        # ====================================================

        self.stdout.write(
            "-" * 165
        )

        self.stdout.write(
            f"Created: {created}   "
            f"Existing: {existing}   "
            f"Skipped: {skipped}   "
            f"Failed: {failed}"
        )

        self.stdout.write(
            f"(blanks={blank_count}, "
            f"dupes={duplicate_count}, "
            f"ignored={ignored_count}, "
            f"invalid={invalid_count})"
        )

        # ====================================================
        # DRY RUN MESSAGE
        # ====================================================

        if DRY_RUN:

            self.stdout.write(
                self.style.WARNING(
                    "\nDRY RUN - no users were created."
                )
            )

            self.stdout.write(
                self.style.WARNING(
                    "Review the output above."
                )
            )

            self.stdout.write(
                self.style.WARNING(
                    "Set DRY_RUN = False and run the "
                    "command again to create the users."
                )
            )

        else:

            self.stdout.write(
                self.style.SUCCESS(
                    f"\nImport completed. "
                    f"{created} user(s) created."
                )
            )